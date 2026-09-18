"""Bounded plan imports and typed, formula-safe production workbook exports."""
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import math
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

if __package__:
    from .control_store import ControlError, ControlStore, PRODUCTION_SCOPE_LABELS, production_scope_context
else:
    from control_store import ControlError, ControlStore, PRODUCTION_SCOPE_LABELS, production_scope_context

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_PLAN_ROWS = 500
PLAN_COLUMNS = ['Xí nghiệp', 'Loại kế hoạch', 'Tháng', 'ID chuyến', 'Chỉ tiêu', 'Sản lượng', 'Số văn bản', 'Ghi chú', 'Quý', 'Năm', 'Từ ngày', 'Đến ngày', 'Tuần']
PLAN_KEYS = ['terminal', 'period_type', 'month', 'voyage_id', 'metric', 'amount', 'reference', 'note', 'quarter', 'year', 'start_date', 'end_date', 'week']
NAVY = '153D39'


def _cell(sheet, row, column, value):
    cell = sheet.cell(row, column)
    if isinstance(value, Decimal):
        value = float(value)
    if isinstance(value, float) and not math.isfinite(value):
        value = None
    cell.value = value
    if isinstance(value, str):
        # An operational name such as '=1+1' is text, never an Excel formula.
        cell.data_type = 's'
    elif isinstance(value, (date, datetime)):
        cell.number_format = 'dd/mm/yyyy hh:mm' if isinstance(value, datetime) else 'dd/mm/yyyy'
    elif isinstance(value, (float, int)):
        cell.number_format = '#,##0.###'
    cell.alignment = Alignment(vertical='top', wrap_text=isinstance(value, str))
    return cell


def _table(sheet, headers, rows, widths=None):
    for col, title in enumerate(headers, 1):
        cell = _cell(sheet, 1, col, title)
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor=NAVY)
    sheet.row_dimensions[1].height = 32
    for row_no, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            cell = _cell(sheet, row_no, col, value)
            cell.font = Font(name='Calibri', size=11)
            if row_no % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F5F3')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}'
    for col in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 22)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'


def _save(book):
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def plan_template():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Kế hoạch'
    _table(sheet, PLAN_COLUMNS, [], {7: 30, 8: 45})
    for column, values in [('A', 'all,cua_lo,ben_thuy'), ('B', 'week,month,quarter,year,custom,voyage'), ('E', 'tonnage,teu')]:
        validation = DataValidation(type='list', formula1=f'"{values}"', allow_blank=False)
        validation.errorTitle = 'Giá trị không hợp lệ'
        validation.error = 'Chọn một giá trị trong danh sách.'
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f'{column}2:{column}{MAX_PLAN_ROWS + 1}')
    for row in sheet.iter_rows(min_row=2, max_row=31, min_col=1, max_col=len(PLAN_COLUMNS)):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='EEF4FC')
            cell.font = Font(name='Calibri', size=11, color='245D8C')
            if cell.column in {3, 4, 7, 9, 11, 12, 13}:
                cell.number_format = '@'
    guide = book.create_sheet('Hướng dẫn')
    _table(guide, ['Trường', 'Cách nhập'], [
        ['Xí nghiệp', 'all: toàn công ty; cua_lo hoặc ben_thuy: từng cảng. Kế hoạch chuyến phải chọn một cảng.'],
        ['Loại kế hoạch', 'week: tuần; month: tháng; quarter: quý; year: năm; custom: khoảng ngày; voyage: toàn chuyến. Chỉ điền các trường của loại kỳ được chọn.'],
        ['Tuần', 'YYYY-Www theo ISO, ví dụ 2026-W38: 14/09–20/09/2026. Tuần đủ thứ Hai đến Chủ nhật, có thể là tuần tương lai. Năm ISO từ 2000 đến 2099; tuần giao năm có thể chứa ngày thuộc năm khác. Chỉ điền với week.'],
        ['Tháng', 'YYYY-MM, chỉ điền với kế hoạch tháng.'],
        ['Quý', 'YYYY-Q1 đến YYYY-Q4, ví dụ 2026-Q3. Chỉ điền với kế hoạch quý.'],
        ['Năm', 'Số nguyên từ 2000 đến 2099, chỉ điền với kế hoạch năm.'],
        ['Từ ngày / Đến ngày', 'YYYY-MM-DD hoặc ô ngày Excel; chỉ điền với custom. Khoảng ngày tối đa 366 ngày, có thể gồm ngày tương lai.'],
        ['ID chuyến', 'ID chuyến từ dashboard, chỉ điền với kế hoạch chuyến. Giữ dạng văn bản.'],
        ['Chỉ tiêu', 'tonnage: tấn thông qua; teu: TEU. Mỗi chỉ tiêu ghi một dòng.'],
        ['Sản lượng', 'Ô số không âm, tối đa 3 chữ số thập phân. Không dùng công thức hoặc số kèm đơn vị.'],
        ['Số văn bản', 'Số hoặc tên văn bản kế hoạch để đối chiếu khi duyệt.'],
        ['Nhập và duyệt', 'Tối đa 500 dòng. Xem trước dữ liệu, nhập bản nháp rồi duyệt trong dashboard.'],
    ], {1: 25, 2: 95})
    return _save(book)


def parse_plan_workbook(content):
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError('Tệp Excel tối đa 2 MB.')
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > 20 * 1024 * 1024:
                raise ValueError('Tệp Excel vượt giới hạn dữ liệu giải nén.')
            if any('vbaproject' in item.filename.lower() or 'externallinks/' in item.filename.lower() for item in entries):
                raise ValueError('Không nhận macro hoặc liên kết tới workbook khác.')
    except zipfile.BadZipFile:
        raise ValueError('Chỉ nhận tệp Excel .xlsx hợp lệ.') from None
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise ValueError('Không đọc được cấu trúc tệp Excel.') from None
    rows, errors = [], []
    try:
        sheet = book['Kế hoạch'] if 'Kế hoạch' in book.sheetnames else book.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_PLAN_ROWS + 1:
            raise ValueError('Mỗi lần nhập tối đa 500 dòng kế hoạch.')
        if sheet.max_column and sheet.max_column > len(PLAN_COLUMNS):
            raise ValueError('Các cột chưa đúng mẫu. Hãy tải mẫu Excel từ dashboard.')
        iterator = sheet.iter_rows(max_row=MAX_PLAN_ROWS + 2, max_col=len(PLAN_COLUMNS))
        first = next(iterator, ())
        header = [cell.value for cell in first]
        while header and header[-1] is None:
            header.pop()
        if header not in [PLAN_COLUMNS, PLAN_KEYS, PLAN_COLUMNS[:12], PLAN_KEYS[:12], PLAN_COLUMNS[:8], PLAN_KEYS[:8]]:
            raise ValueError('Các cột chưa đúng mẫu. Hãy tải mẫu Excel từ dashboard.')
        seen = set()
        for row_no, cells in enumerate(iterator, 2):
            if not any(cell.value is not None for cell in cells):
                continue
            if row_no > MAX_PLAN_ROWS + 1:
                raise ValueError('Mỗi lần nhập tối đa 500 dòng kế hoạch.')
            if any(cell.value is not None for cell in cells[len(header):]):
                errors.append({'row': row_no, 'message': 'Có dữ liệu ngoài các cột của mẫu. Hãy tải mẫu Excel mới từ dashboard.'})
                continue
            if any(cell.data_type == 'f' for cell in cells):
                errors.append({'row': row_no, 'message': 'Chuyển công thức thành giá trị trước khi nhập.'})
                continue
            record = dict(zip(PLAN_KEYS, [cell.value for cell in cells]))
            for key in ['terminal', 'period_type', 'month', 'quarter', 'week', 'metric', 'reference', 'note']:
                record[key] = str(record[key] or '').strip()
            try:
                amount = record['amount']
                if isinstance(amount, bool) or not isinstance(amount, (int, float, Decimal)):
                    raise ValueError('Sản lượng phải là ô số, không kèm đơn vị hoặc dấu phân cách dạng văn bản.')
                amount = Decimal(str(amount))
                if not amount.is_finite() or amount < 0 or amount > Decimal('1000000000000') or amount != amount.quantize(Decimal('.001')):
                    raise ValueError('Sản lượng phải không âm, tối đa 3 chữ số thập phân.')
                record['amount'] = str(amount)
                if record['terminal'] not in {'all', 'cua_lo', 'ben_thuy'} or record['metric'] not in {'tonnage', 'teu'}:
                    raise ValueError('Xí nghiệp hoặc chỉ tiêu chưa đúng danh mục trong mẫu.')
                for field in ['month', 'quarter', 'year', 'voyage_id', 'start_date', 'end_date', 'week']:
                    if record[field] == '':
                        record[field] = None
                if record['voyage_id'] is not None:
                    value = Decimal(str(record['voyage_id']))
                    if value != value.to_integral_value() or not 1 <= value <= 2147483647:
                        raise ValueError('Kế hoạch chuyến cần ID nguyên dương.')
                    record['voyage_id'] = int(value)
                if record['year'] is not None:
                    value = Decimal(str(record['year']))
                    if value != value.to_integral_value():
                        raise ValueError('Năm kế hoạch phải là số nguyên.')
                    record['year'] = int(value)
                for field in ['start_date', 'end_date']:
                    if isinstance(record[field], datetime):
                        record[field] = record[field].date().isoformat()
                    elif isinstance(record[field], date):
                        record[field] = record[field].isoformat()
                validated = ControlStore.validate_plan(record)
                if len(record['reference']) > 300 or len(record['note']) > 2000:
                    raise ValueError('Số văn bản hoặc ghi chú quá dài.')
                key = tuple(validated[field] for field in ['terminal', 'period_type', 'period_key', 'metric'])
                if key in seen:
                    raise ValueError('Trùng xí nghiệp, kỳ/chuyến và chỉ tiêu trong cùng tệp.')
                seen.add(key)
                rows.append(record)
            except ControlError as exc:
                errors.append({'row': row_no, 'message': exc.message})
            except (ValueError, TypeError, InvalidOperation) as exc:
                errors.append({'row': row_no, 'message': str(exc) if isinstance(exc, ValueError) else 'Số liệu hoặc kỳ kế hoạch không hợp lệ.'})
    finally:
        book.close()
    if not rows and not errors:
        errors.append({'row': 2, 'message': 'Tệp chưa có dòng kế hoạch.'})
    return {'rows': rows, 'errors': errors, 'valid': bool(rows) and not errors}


def report_workbook(report, operations, *, title='Báo cáo sản lượng', shifts=None, planning=None):
    book = Workbook()
    summary_sheet = book.active
    summary_sheet.title = 'Tổng hợp'
    meta = report.get('meta', {})
    filters = meta.get('filters', {})
    overview = report.get('overview') or report.get('summary', {})
    tonnage = overview.get('total_tonnage', overview.get('tonnage'))
    teu = overview.get('total_teu', overview.get('teu'))
    scope = production_scope_context(report)
    rows = [
        ['Báo cáo', title, None], ['Mã phiên dữ liệu', meta.get('report_id'), None],
        ['Từ ngày', date.fromisoformat(filters['start_date']) if filters.get('start_date') else None, None],
        ['Đến ngày', date.fromisoformat(filters['end_date']) if filters.get('end_date') else None, None],
        ['Xí nghiệp', filters.get('terminal'), None],
        ['Đọc nguồn lúc', meta.get('source_read_at', meta.get('generated_at')), None],
        ['Sản lượng thông qua', tonnage, 'Tấn'], ['Container tác nghiệp', teu, 'TEU'],
        ['Số dòng nguồn', overview.get('record_count'), 'Dòng'],
        ['Trạng thái số liệu tấn', overview.get('tonnage_status'), None],
        ['Trạng thái số liệu TEU', overview.get('teu_status'), None],
        ['Bộ lọc dòng tác nghiệp', meta.get('operations_filter', 'all'), None],
        ['Số dòng trong tệp', len(operations), 'Dòng'],
        ['Phạm vi sản lượng', scope['production_scope_label'], None],
        ['Mã phạm vi sản lượng', scope['production_scope'] or 'legacy', None],
        ['Phiên bản quy tắc cầu cập đầu tiên', scope['berth_rule_version'] or 'Chưa được lưu trong bản dữ liệu này', None],
    ]
    if scope['legacy_scope']:
        rows.append(['Phân loại theo cầu cập đầu tiên', 'Bản dữ liệu cũ chưa lưu phạm vi này; không tự áp dụng quy tắc hiện tại cho lịch sử.', None])
    planning = planning if planning is not None else report.get('planning')
    if planning is not None:
        rows.append(['Kế hoạch tại thời điểm chốt', 'Đã lưu' if planning.get('captured') else 'Chưa được lưu trong bản chốt này', None])
        rows.append(['Thời điểm lưu kế hoạch', planning.get('captured_at'), None])
        if planning.get('reason'):
            rows.append(['Phạm vi đối chiếu kế hoạch', planning['reason'], None])
        if planning.get('eligible'):
            rows.append(['Phạm vi kế hoạch đã chốt', scope['production_scope_label'], None])
    for field, value in meta.get('selection', {}).items():
        if value is not None and field not in {'operation_filter', 'production_scope'}:
            labels = {'day': 'Ngày tác nghiệp được chọn', 'terminal': 'Xí nghiệp được chọn', 'cargo': 'Nhóm hàng được chọn', 'customer_id': 'ID khách hàng', 'customer_terminal': 'Xí nghiệp khách hàng', 'voyage_id': 'ID chuyến', 'issue': 'Nhóm đối soát'}
            rows.append([labels.get(field, field), str(value) if field in {'customer_id', 'voyage_id'} else value, None])
    _table(summary_sheet, ['Nội dung', 'Giá trị', 'Đơn vị'], rows, {1: 34, 2: 48, 3: 18})
    sheet = book.create_sheet('Tác nghiệp')
    columns = [
        ('terminal_name', 'Xí nghiệp'), ('id', 'ID tác nghiệp'), ('operation_code', 'Mã tác nghiệp'),
        ('operation_date', 'Ngày tác nghiệp'), ('shift_code', 'Ca'), ('vessel_name', 'Tàu'),
        ('voyage_code', 'Mã chuyến'), ('cargo_name', 'Hàng hóa'), ('customer_name', 'Khách hàng'),
        ('direction', 'Hướng hàng'), ('quantity', 'Số lượng nguồn'), ('quantity_unit_name', 'Đơn vị số lượng'),
        ('weight', 'Trọng lượng nguồn'), ('weight_unit_name', 'Đơn vị trọng lượng'), ('tonnage', 'Tấn'), ('teu', 'TEU'),
        ('production_scope', 'Phạm vi sản lượng'), ('berth_assignment_status', 'Trạng thái xác định cầu đầu tiên'),
        ('initial_berth_id', 'ID cầu cập đầu tiên'), ('initial_berth_code', 'Mã cầu cập đầu tiên'),
        ('initial_berth_at', 'Thời điểm cập cầu đầu tiên'),
    ]
    values = []
    berth_status_labels = {'assigned': 'Đã xác định', 'missing': 'Thiếu dữ liệu cầu', 'ambiguous': 'Không xác định duy nhất'}
    for row in operations:
        exported = []
        for key, _ in columns:
            value = row.get(key)
            if key == 'operation_date' and value:
                value = date.fromisoformat(value)
            elif key in {'id', 'operation_code', 'voyage_code', 'initial_berth_id', 'initial_berth_code'} and value is not None:
                value = str(value)
            elif key == 'initial_berth_at' and value is not None:
                value = value.isoformat() if isinstance(value, (date, datetime)) else str(value)
            elif key == 'production_scope':
                value = scope['production_scope_label'] if value is None else PRODUCTION_SCOPE_LABELS.get(value, str(value))
            elif key == 'berth_assignment_status':
                value = berth_status_labels.get(value, value or 'Chưa lưu bằng chứng phân loại')
            exported.append(value)
        values.append(exported)
    _table(sheet, [label for _, label in columns], values, {1: 20, 2: 18, 3: 20, 4: 18, 8: 27, 9: 28, 17: 24, 18: 32, 19: 22, 20: 24, 21: 32})
    if shifts:
        sheet = book.create_sheet('Theo ca')
        _table(sheet, ['Ngày', 'Ca', 'Xí nghiệp', 'Tấn', 'TEU', 'Dòng tác nghiệp'], [
            [row.get('date') or row.get('operation_date'), row.get('shift_code'), row.get('terminal_name'), row.get('tonnage'), row.get('teu'), row.get('record_count')]
            for row in shifts
        ])
    if planning is not None:
        sheet = book.create_sheet('Kế hoạch đã chốt')
        if not planning.get('captured') or not planning.get('eligible'):
            _table(sheet, ['Nội dung', 'Giá trị'], [
                ['Trạng thái', planning.get('reason') or 'Bản chốt chưa lưu đối chiếu kế hoạch.'],
            ], {1: 28, 2: 100})
        else:
            labels = {'missing_plan': 'Chưa có kế hoạch đã duyệt', 'incomplete_actual': 'Số liệu thực hiện chưa đầy đủ',
                      'zero_target': 'Chỉ tiêu bằng 0', 'ready': 'Đủ dữ liệu đối chiếu'}
            actual_labels = {'ready': 'Đầy đủ', 'empty': 'Không phát sinh', 'partial': 'Chưa đầy đủ', 'unavailable': 'Chưa có dữ liệu'}
            _table(sheet, ['Xí nghiệp', 'Tháng kế hoạch', 'Chỉ tiêu', 'ID kế hoạch', 'Phiên bản kế hoạch', 'Kế hoạch',
                           'Thực hiện trong kỳ', 'Còn lại', '% hoàn thành', 'Trạng thái thực hiện', 'Trạng thái đối chiếu',
                           'Căn cứ duyệt', 'ID người duyệt', 'Thời điểm duyệt'], [
                [row.get('terminal_name'), planning.get('period', {}).get('month'), 'Tấn' if row['metric'] == 'tonnage' else 'TEU',
                 str(row['plan_id']) if row.get('plan_id') is not None else None, row.get('plan_version'), row.get('target'),
                 row.get('actual'), row.get('remaining'), row.get('completion_percent'),
                 actual_labels.get(row.get('actual_status'), row.get('actual_status')), labels.get(row.get('status'), row.get('status')),
                 row.get('reference'), str(row['approved_by']) if row.get('approved_by') is not None else None, row.get('approved_at')]
                for row in planning.get('rows', [])
            ], {1: 20, 2: 20, 10: 28, 11: 35, 12: 50, 14: 32})
    throughput = report.get('throughput_progress')
    if throughput is not None:
        sheet = book.create_sheet('Mục tiêu thông qua')
        labels = {'week': 'Tuần', 'month': 'Tháng', 'quarter': 'Quý', 'year': 'Năm', 'custom': 'Khoảng ngày'}
        if throughput.get('items'):
            _table(sheet, ['Loại kỳ', 'Kỳ kế hoạch', 'Từ ngày', 'Đến ngày', 'Mục tiêu tấn', 'Thực hiện tấn',
                           '% hoàn thành xác nhận', '% tạm tính', 'Đã đạt', 'Trạng thái', 'Nguồn mục tiêu',
                           'ID / phiên bản kế hoạch', 'Căn cứ duyệt'], [
                [labels.get(item['period_type'], item['period_type']), item['period_key'], item['start_date'], item['end_date'],
                 item['target'], item['actual'], item['completion_percent'], item['provisional_completion_percent'],
                 'Đã đạt' if item['achieved'] else 'Chưa xác nhận đạt' if item['provisional'] else 'Chưa đạt',
                 item.get('reason') or item['status'], {'company': 'Toàn công ty', 'terminals': 'Tổng hai cảng', 'terminal': 'Cảng được chọn'}.get(item['target_source']),
                 '; '.join(f"{plan['id']} / v{plan['version']}" for plan in item['plans']),
                 '; '.join(plan['reference'] for plan in item['plans'])]
                for item in throughput['items']
            ], {2: 30, 7: 28, 8: 22, 10: 44, 12: 30, 13: 50})
        else:
            _table(sheet, ['Nội dung', 'Giá trị'], [['Trạng thái', throughput.get('reason')]], {1: 25, 2: 100})
    return _save(book)
