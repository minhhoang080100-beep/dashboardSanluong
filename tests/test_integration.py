"""Real auth and control storage with synthetic source facts; never reads TOS."""
from copy import deepcopy
from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from backend import main, repository
from backend.control_store import ControlStore
from backend.reporting import ReportingService
from backend.workbook_io import plan_template
from test_reporting import StubRepository, fact

PASSWORD = 'Synthetic-Only-Password-2026'


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(repository, 'vietnam_today', lambda: date(2026, 9, 13))
    store = ControlStore(tmp_path / 'control.sqlite3')
    first = store.bootstrap_admin()
    initial = store.login('admin', first['temporary_password'])
    admin_session = store.change_password(initial['token'], first['temporary_password'], PASSWORD)
    admin = admin_session['user']
    viewer = store.create_user(admin, username='viewer', display_name='Synthetic viewer', role='viewer', terminals=['cua_lo'])
    initial_viewer = store.login('viewer', viewer['temporary_password'])
    viewer_session = store.change_password(initial_viewer['token'], viewer['temporary_password'], PASSWORD)
    rows = [fact(1, day=1, weight='50'), fact(2, day=11, weight='10'), fact(3, day=12, weight='20'),
            fact(4, day=12, voyage=None, weight='7'), fact(1, day=12, terminal='ben_thuy', weight='5')]
    repo = StubRepository(rows)
    def valid_voyages(pairs):
        if repo.failure:
            from backend.database import DatabaseUnavailable
            raise DatabaseUnavailable()
        existing = {(row['terminal_id'], str(row['vessel_id'])) for row in repo.rows if row['vessel_id'] is not None}
        return set(pairs) & existing
    repo.validate_voyages = valid_voyages
    repo.search_voyages = lambda terminal, search='', limit=30: [
        {'terminal': terminal, 'voyage_id': '101', 'vessel_name': 'Synthetic future vessel',
         'voyage_code': 'SYNTHETIC-101', 'arrival_date': None, 'departure_date': None}]
    service = ReportingService(repo)
    main.app.state.control_store = store
    main.app.state.reporting_service = service
    try:
        with TestClient(main.app) as client:
            yield client, store, repo, service, {'Authorization': 'Bearer ' + admin_session['token']}, {'Authorization': 'Bearer ' + viewer_session['token']}
    finally:
        del main.app.state.control_store
        del main.app.state.reporting_service


def get_report(api, terminal='all', start='2026-09-11', end='2026-09-12', viewer=False, refresh=False):
    client, _, _, _, admin, limited = api
    response = client.get('/api/dashboard', headers=limited if viewer else admin,
                          params={'start_date': start, 'end_date': end, 'terminal': terminal, 'refresh': refresh})
    assert response.status_code == 200, response.text
    return response.json()


def test_every_production_path_requires_auth_and_liveness_does_not_read_source(api):
    client, _, repo, _, _, _ = api
    for path in ['/api/dashboard', '/api/overview', '/api/top-customers', '/api/voyages/cua_lo/101',
                 '/api/reports/missing/operations', '/api/reports/missing/export.xlsx', '/api/closed-reports',
                 '/api/admin/metrics', '/api/health', '/api/voyages/cua_lo/101/progress']:
        assert client.get(path).status_code == 401, path
    assert client.get('/api/health/live').json() == {'status': 'ok'}
    assert repo.calls == []


def test_scope_is_checked_before_cache_drilldown_export_and_progress(api):
    client, _, repo, _, _, viewer = api
    report = get_report(api)
    rid = report['meta']['report_id']
    calls = len(repo.calls)
    for path in ['/api/dashboard?terminal=all', '/api/dashboard?terminal=ben_thuy',
                 f'/api/reports/{rid}/operations?terminal=cua_lo', f'/api/reports/{rid}/export.xlsx?terminal=cua_lo',
                 '/api/voyages/ben_thuy/101/progress', '/api/admin/metrics', '/api/health']:
        assert client.get(path, headers=viewer).status_code == 403, path
    assert len(repo.calls) == calls
    allowed = get_report(api, terminal='cua_lo', viewer=True)
    assert {row['terminal_id'] for row in allowed['voyages']} == {'cua_lo'}
    assert client.get(f"/api/reports/{allowed['meta']['report_id']}/operations?terminal=ben_thuy", headers=viewer).status_code == 403
    assert client.post('/api/closed-reports', headers=viewer, json={'report_id': allowed['meta']['report_id']}).status_code == 403


def test_snapshot_drilldown_excel_and_voyage_totals_match_original_source_read(api):
    client, _, repo, _, admin, _ = api
    report = get_report(api)
    rid = report['meta']['report_id']
    repo.rows[2]['native_weight'] = 999  # Source later changes; current report stays fixed.
    detail = client.get(f'/api/reports/{rid}/operations', headers=admin).json()
    assert detail['summary']['tonnage'] == report['overview']['total_tonnage'] == 42
    assert detail['operations']['total'] == 4
    assert len(repo.calls) == 1
    voyage = client.get('/api/voyages/cua_lo/101', headers=admin, params={
        'start_date': '2026-09-11', 'end_date': '2026-09-12', 'report_id': rid,
        'operation_filter': 'with_values'}).json()
    assert voyage['summary']['tonnage'] == 30
    assert voyage['meta']['report_id'] == rid
    wrong_period = client.get('/api/voyages/cua_lo/101', headers=admin, params={
        'start_date': '2026-09-01', 'end_date': '2026-09-12', 'report_id': rid})
    assert wrong_period.status_code == 422
    response = client.get(f'/api/reports/{rid}/export.xlsx', headers=admin)
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    book = load_workbook(BytesIO(response.content), data_only=False)
    assert book['Tác nghiệp'].max_row == 5
    assert sum(row[14].value for row in list(book['Tác nghiệp'].rows)[1:]) == 42
    assert len(repo.calls) == 1


def test_closed_reports_are_trusted_immutable_and_compare_only_same_scope(api):
    client, _, repo, _, admin, viewer = api
    report = get_report(api)
    rid = report['meta']['report_id']
    fake = client.post('/api/closed-reports', headers=admin, json={'report_id': rid, 'overview': {'total_tonnage': 0}})
    assert fake.status_code == 422
    closed = client.post('/api/closed-reports', headers=admin, json={'report_id': rid, 'title': 'Synthetic closure'})
    assert closed.status_code == 201
    identifier = closed.json()['id']
    compare_path = f'/api/closed-reports/{identifier}/compare'
    assert client.post(compare_path, headers=admin, json={'report_id': rid}).json()['changed'] is False
    repo.rows[2]['native_weight'] = 25
    current = get_report(api, refresh=True)
    comparison = client.post(compare_path, headers=admin, json={'report_id': current['meta']['report_id']}).json()
    assert comparison['changed'] is True
    assert next(row for row in comparison['changes'] if row['metric'] == 'total_tonnage')['delta'] == 5
    saved = client.get(f'/api/closed-reports/{identifier}', headers=admin).json()
    assert saved['report']['overview']['total_tonnage'] == 42
    assert client.get(f'/api/closed-reports/{identifier}/export.xlsx', headers=admin).status_code == 200
    assert client.get(f'/api/closed-reports/{identifier}', headers=viewer).status_code == 403
    different = get_report(api, start='2026-09-01')
    assert client.post(compare_path, headers=admin, json={'report_id': different['meta']['report_id']}).status_code == 422


def test_plan_progress_requires_approved_matching_period_and_whole_voyage_scope(api):
    client, _, _, _, admin, _ = api
    monthly = {'terminal': 'cua_lo', 'period_type': 'month', 'month': '2026-09', 'metric': 'tonnage', 'amount': 100, 'reference': 'Synthetic approved plan'}
    plan = client.post('/api/plans', headers=admin, json=monthly).json()
    report = get_report(api, terminal='cua_lo', start='2026-09-01')
    progress_path = f"/api/reports/{report['meta']['report_id']}/plan-progress"
    progress = client.get(progress_path, headers=admin).json()
    assert progress['eligible'] and progress['rows'][0]['status'] == 'missing_plan'
    assert client.post(f"/api/plans/{plan['id']}/approve", headers=admin, json={'expected_revision': 1}).status_code == 200
    progress = client.get(progress_path, headers=admin).json()
    tonne_row = next(row for row in progress['rows'] if row['metric'] == 'tonnage')
    assert tonne_row['target'] == 100 and tonne_row['actual'] == 87 and tonne_row['remaining'] == 13
    partial_period = get_report(api, terminal='cua_lo')
    assert client.get(f"/api/reports/{partial_period['meta']['report_id']}/plan-progress", headers=admin).json()['eligible'] is False
    voyage_plan = {**monthly, 'period_type': 'voyage', 'month': None, 'voyage_id': 101}
    created = client.post('/api/plans', headers=admin, json=voyage_plan).json()
    client.post(f"/api/plans/{created['id']}/approve", headers=admin, json={'expected_revision': 1})
    voyage = client.get('/api/voyages/cua_lo/101/progress', headers=admin).json()
    assert voyage['summary']['tonnage'] == 80  # Includes Sep1; period detail only has 30.
    assert voyage['planning'][0]['remaining'] == 20
    assert voyage['productivity'] is None and voyage['estimated_completion'] is None


def test_plan_excel_preview_does_not_write_and_rejects_foreign_terminal(api):
    client, store, _, _, admin, viewer = api
    book = load_workbook(BytesIO(plan_template()))
    sheet = book['Kế hoạch']
    for col, value in enumerate(['ben_thuy', 'month', '2026-09', None, 'teu', 100, 'Synthetic document', ''], 1):
        sheet.cell(2, col, value)
    output = BytesIO(); book.save(output)
    response = client.post('/api/plans/import/preview', headers=admin, files={'file': ('plans.xlsx', output.getvalue())})
    assert response.status_code == 200 and response.json()['valid']
    assert client.get('/api/plans', headers=admin).json()['total'] == 0
    assert client.post('/api/plans/import/preview', headers=viewer, files={'file': ('plans.xlsx', output.getvalue())}).status_code == 403
    imported = client.post('/api/plans/import', headers=admin, json={'rows': response.json()['rows']})
    assert imported.status_code == 201 and imported.json()['items'][0]['status'] == 'draft'


def test_auth_cors_and_expired_report_are_explicit(api):
    client, _, _, _, admin, _ = api
    origin = 'https://dashboard-sanluong.vercel.app'
    preflight = client.options('/api/plans', headers={'Origin': origin, 'Access-Control-Request-Method': 'POST', 'Access-Control-Request-Headers': 'Authorization,Content-Type'})
    assert preflight.status_code == 200 and preflight.headers['access-control-allow-origin'] == origin
    assert client.get('/api/dashboard', headers={'Origin': origin}).headers['access-control-allow-origin'] == origin
    expired = client.get('/api/reports/nonexistent/operations', headers=admin)
    assert expired.status_code == 410 and expired.json()['detail']['code'] == 'REPORT_EXPIRED'


def test_planning_lookup_and_all_mutations_validate_terminal_voyage_identity(api):
    client, store, repo, _, admin, viewer = api
    assert client.get('/api/planning/voyages?terminal=cua_lo', headers=viewer).status_code == 403
    lookup = client.get('/api/planning/voyages?terminal=cua_lo&search=Synthetic', headers=admin)
    assert lookup.status_code == 200 and lookup.json()['items'][0]['arrival_date'] is None
    plan = {'terminal': 'cua_lo', 'period_type': 'voyage', 'voyage_id': 999,
            'metric': 'tonnage', 'amount': 100, 'reference': 'Approved document'}
    assert client.post('/api/plans', headers=admin, json=plan).json()['detail']['code'] == 'PLAN_VOYAGE_NOT_FOUND'
    monthly = {**plan, 'period_type': 'month', 'voyage_id': None, 'month': '2026-09'}
    rejected = client.post('/api/plans/import', headers=admin, json={'rows': [monthly, plan]})
    assert rejected.status_code == 422 and store.list_plans(store.authenticate(admin['Authorization'].split()[1]))['total'] == 0
    plan['voyage_id'] = 101
    created = client.post('/api/plans', headers=admin, json=plan).json()
    path = f"/api/plans/{created['id']}"
    assert client.patch(path, headers=admin, json={'expected_revision': 1, 'voyage_id': 999}).status_code == 422
    assert client.get(path, headers=admin).json()['revision'] == 1
    # A voyage removed from the source between drafting and approval cannot be approved.
    repo.rows[:] = [row for row in repo.rows if row['terminal_id'] != 'cua_lo']
    assert client.post(path + '/approve', headers=admin, json={'expected_revision': 1}).status_code == 422
    assert client.get(path, headers=admin).json()['status'] == 'draft'
    assert client.post(path + '/cancel', headers=admin, json={'expected_revision': 1, 'note': 'Wrong voyage'}).status_code == 200


def test_monthly_planning_and_cancellation_work_during_sql_outage(api):
    client, _, repo, _, admin, _ = api
    repo.failure = True
    monthly = {'terminal': 'cua_lo', 'period_type': 'month', 'month': '2026-09',
               'metric': 'tonnage', 'amount': 100, 'reference': 'Approved document'}
    created = client.post('/api/plans', headers=admin, json=monthly)
    assert created.status_code == 201
    path = f"/api/plans/{created.json()['id']}"
    changed = client.patch(path, headers=admin, json={'expected_revision': 1, 'amount': 150})
    assert changed.status_code == 200 and changed.json()['revision'] == 2
    stale = client.post(path + '/approve', headers=admin, json={'expected_revision': 1})
    assert stale.status_code == 409 and stale.json()['detail']['code'] == 'PLAN_CONFLICT'
    assert client.post(path + '/approve', headers=admin, json={'expected_revision': 2}).status_code == 200
    assert client.get('/api/plans', headers=admin).status_code == 200
    assert client.get('/api/plans/template.xlsx', headers=admin).status_code == 200
    book = load_workbook(BytesIO(plan_template()))
    for col, value in enumerate(['cua_lo', 'month', '2026-09', None, 'teu', 10, 'Monthly document', ''], 1):
        book['Kế hoạch'].cell(2, col, value)
    output = BytesIO(); book.save(output)
    preview = client.post('/api/plans/import/preview', headers=admin, files={'file': ('plans.xlsx', output.getvalue())})
    assert preview.status_code == 200 and preview.json()['valid']
    voyage = {**monthly, 'period_type': 'voyage', 'month': None, 'voyage_id': 101}
    assert client.post('/api/plans', headers=admin, json=voyage).status_code == 503
    assert client.get('/api/plans', headers=admin).json()['total'] == 1


def test_excel_preview_rejects_unknown_or_mismatched_voyage_without_writing(api):
    client, _, _, _, admin, _ = api
    book = load_workbook(BytesIO(plan_template()))
    for col, value in enumerate(['ben_thuy', 'voyage', None, 999, 'tonnage', 100, 'Document', ''], 1):
        book['Kế hoạch'].cell(2, col, value)
    output = BytesIO(); book.save(output)
    response = client.post('/api/plans/import/preview', headers=admin, files={'file': ('plans.xlsx', output.getvalue())})
    assert response.status_code == 422 and response.json()['detail']['code'] == 'PLAN_VOYAGE_NOT_FOUND'
    assert client.get('/api/plans', headers=admin).json()['total'] == 0


def test_closed_reports_freeze_approved_plan_version_actuals_and_excel(api):
    client, _, _, _, admin, _ = api
    body = {'terminal': 'cua_lo', 'month': '2026-09', 'metric': 'tonnage', 'amount': 100, 'reference': '=Original document'}
    first = client.post('/api/plans', headers=admin, json=body).json()
    client.post(f"/api/plans/{first['id']}/approve", headers=admin, json={'expected_revision': 1})
    report = get_report(api, terminal='cua_lo', start='2026-09-01')
    closed = client.post('/api/closed-reports', headers=admin, json={'report_id': report['meta']['report_id']}).json()
    revised = client.post('/api/plans', headers=admin, json={**body, 'amount': 200, 'reference': 'Revised'}).json()
    client.post(f"/api/plans/{revised['id']}/approve", headers=admin, json={'expected_revision': 1})
    saved = client.get(f"/api/closed-reports/{closed['id']}", headers=admin).json()
    planning = saved['planning']
    assert planning == saved['report']['planning'] and planning['captured'] and planning['eligible']
    row = next(item for item in planning['rows'] if item['metric'] == 'tonnage')
    assert (row['plan_id'], row['plan_version'], row['target'], row['actual'], row['completion_percent']) == (first['id'], 1, 100, 87, 87)
    assert row['reference'] == '=Original document' and row['approved_at']
    response = client.get(f"/api/closed-reports/{closed['id']}/export.xlsx", headers=admin)
    book = load_workbook(BytesIO(response.content), data_only=False)
    sheet = book['Kế hoạch đã chốt']
    assert sheet['F2'].value == 100 and sheet['G2'].value == 87 and sheet['I2'].value == 87
    assert sheet['L2'].value == '=Original document' and sheet['L2'].data_type == 's'
    book.close()


def test_closed_report_partial_period_explains_no_monthly_comparison(api):
    client, _, _, _, admin, _ = api
    report = get_report(api)
    closed = client.post('/api/closed-reports', headers=admin, json={'report_id': report['meta']['report_id']}).json()
    planning = client.get(f"/api/closed-reports/{closed['id']}", headers=admin).json()['planning']
    assert planning['captured'] and not planning['eligible'] and planning['rows'] == [] and planning['reason']


def test_old_closed_report_does_not_backfill_current_plans(api):
    client, store, _, _, admin, _ = api
    user = store.authenticate(admin['Authorization'].split()[1])
    closed = store.close_report(user, 'cua_lo', '2026-09-01', '2026-09-13', {'overview': {}}, [])
    plan = store.create_plan(user, terminal='cua_lo', month='2026-09', metric='tonnage', amount=100, reference='New document')
    store.approve_plan(user, plan['id'])
    saved = client.get(f"/api/closed-reports/{closed['id']}", headers=admin).json()
    assert not saved['planning']['captured'] and saved['planning']['rows'] == []
    assert 'planning' not in saved['report']
    response = client.get(f"/api/closed-reports/{closed['id']}/export.xlsx", headers=admin)
    book = load_workbook(BytesIO(response.content))
    assert 'Không sử dụng kế hoạch hiện tại' in book['Kế hoạch đã chốt']['B2'].value
    book.close()


@pytest.mark.parametrize('action', ['approve', 'update'])
def test_source_validation_cannot_approve_or_overwrite_concurrently_changed_draft(api, monkeypatch, action):
    client, store, repo, _, admin, _ = api
    body = {'terminal': 'cua_lo', 'period_type': 'voyage', 'voyage_id': 101, 'metric': 'tonnage', 'amount': 100, 'reference': 'Document'}
    plan = client.post('/api/plans', headers=admin, json=body).json()
    user = store.authenticate(admin['Authorization'].split()[1])
    original = repo.validate_voyages
    def change_during_validation(pairs):
        store.update_plan(user, plan['id'], 1, amount=200)
        return original(pairs)
    monkeypatch.setattr(repo, 'validate_voyages', change_during_validation)
    path = f"/api/plans/{plan['id']}"
    response = client.post(path + '/approve', headers=admin, json={'expected_revision': 1}) if action == 'approve' else client.patch(
        path, headers=admin, json={'expected_revision': 1, 'amount': 300})
    assert response.status_code == 409 and response.json()['detail']['code'] == 'PLAN_CONFLICT'
    saved = client.get(path, headers=admin).json()
    assert saved['status'] == 'draft' and saved['amount'] == 200 and saved['revision'] == 2
