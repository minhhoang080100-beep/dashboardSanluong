"""Excel regressions use synthetic data and inspect actual saved XLSX files."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import zipfile

from openpyxl import Workbook, load_workbook
import pytest

from backend import workbook_io


LEGACY_PLAN_COLUMNS_8 = ['Xí nghiệp', 'Loại kế hoạch', 'Tháng', 'ID chuyến', 'Chỉ tiêu', 'Sản lượng', 'Số văn bản', 'Ghi chú']
LEGACY_PLAN_KEYS_8 = ['terminal', 'period_type', 'month', 'voyage_id', 'metric', 'amount', 'reference', 'note']
LEGACY_PLAN_COLUMNS_12 = [*LEGACY_PLAN_COLUMNS_8, 'Quý', 'Năm', 'Từ ngày', 'Đến ngày']
LEGACY_PLAN_KEYS_12 = [*LEGACY_PLAN_KEYS_8, 'quarter', 'year', 'start_date', 'end_date']


def make_plan(rows=None, headers=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Kế hoạch"
    sheet.append(headers if headers is not None else workbook_io.PLAN_COLUMNS)
    for row in rows or []:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def month_row(**changes):
    values = {"terminal": "cua_lo", "period_type": "month", "month": "2026-09", "voyage_id": None,
              "metric": "tonnage", "amount": 120.125, "reference": "PLAN-EXAMPLE", "note": None, **changes}
    return [values.get(key) for key in workbook_io.PLAN_KEYS]


def replace_zip_member(content, name, replacement):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as original, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as changed:
        for member in original.infolist():
            changed.writestr(member.filename, replacement(original.read(member.filename)) if member.filename == name else original.read(member.filename))
    return output.getvalue()


def test_template_roundtrips_headers_validation_and_empty_state(tmp_path):
    target = tmp_path / "plan-template.xlsx"
    target.write_bytes(workbook_io.plan_template())
    book = load_workbook(target)
    assert book.sheetnames == ["Kế hoạch", "Hướng dẫn"]
    assert [cell.value for cell in book["Kế hoạch"][1]] == workbook_io.PLAN_COLUMNS
    assert workbook_io.PLAN_COLUMNS == [*LEGACY_PLAN_COLUMNS_12, 'Tuần']
    assert workbook_io.PLAN_KEYS == [*LEGACY_PLAN_KEYS_12, 'week']
    assert len(book["Kế hoạch"].data_validations.dataValidation) == 3
    assert book["Kế hoạch"]["D2"].number_format == "@"
    assert book["Kế hoạch"]["M2"].number_format == "@"
    period_validation = next(item for item in book['Kế hoạch'].data_validations.dataValidation if 'B2' in item)
    assert period_validation.formula1 == '"week,month,quarter,year,custom,voyage"'
    book.close()
    preview = workbook_io.parse_plan_workbook(target.read_bytes())
    assert preview["valid"] is False and preview["rows"] == []
    assert preview["errors"][0]["row"] == 2


@pytest.mark.parametrize("headers", [workbook_io.PLAN_COLUMNS, workbook_io.PLAN_KEYS])
def test_import_accepts_typed_numbers_text_identifiers_and_null_optional_fields(headers):
    voyage = month_row(period_type="voyage", month=None, voyage_id="00000101", metric="teu", amount=0, reference="000012")
    preview = workbook_io.parse_plan_workbook(make_plan([month_row(), voyage], headers))
    assert preview["valid"] and not preview["errors"]
    monthly, ship = preview["rows"]
    assert monthly["amount"] == "120.125" and monthly["voyage_id"] is None
    assert monthly["note"] == ""
    assert ship["voyage_id"] == 101 and ship["month"] is None
    assert ship["amount"] == "0" and ship["reference"] == "000012"


@pytest.mark.parametrize("field,value", [("amount", "=SUM(1,2)"), ("reference", '=HYPERLINK("https://example.invalid","test")'), ("note", "=1+1")])
def test_import_rejects_formula_cells_without_evaluating_them(field, value):
    preview = workbook_io.parse_plan_workbook(make_plan([month_row(**{field: value})]))
    assert not preview["valid"] and preview["rows"] == []
    assert preview["errors"][0]["row"] == 2
    assert "công thức" in preview["errors"][0]["message"]


@pytest.mark.parametrize("value", ["120.125", "1,200", "12 tấn", True, None, -1, 0.0001, 1000000000001])
def test_import_requires_bounded_nonnegative_native_number_cells(value):
    preview = workbook_io.parse_plan_workbook(make_plan([month_row(amount=value)]))
    assert not preview["valid"] and preview["rows"] == []
    assert preview["errors"][0]["row"] == 2


@pytest.mark.parametrize("headers", [list(reversed(workbook_io.PLAN_COLUMNS)), ["Wrong", *workbook_io.PLAN_COLUMNS[1:]],
                                     LEGACY_PLAN_COLUMNS_12[:-1], [*workbook_io.PLAN_COLUMNS, 'Extra'],
                                     [*LEGACY_PLAN_COLUMNS_12, 'Week'], [*LEGACY_PLAN_COLUMNS_12, 'week']])
def test_wrong_headers_are_rejected(headers):
    with pytest.raises(ValueError, match="cột"):
        workbook_io.parse_plan_workbook(make_plan([month_row()], headers))


@pytest.mark.parametrize('headers', [LEGACY_PLAN_COLUMNS_8, LEGACY_PLAN_KEYS_8, LEGACY_PLAN_COLUMNS_12, LEGACY_PLAN_KEYS_12])
def test_legacy_templates_still_import_without_moving_existing_columns(headers):
    rows = [month_row()[:len(headers)], month_row(period_type='voyage', month=None, voyage_id='00101')[:len(headers)]]
    if len(headers) == 12:
        rows.append(month_row(period_type='quarter', month=None, quarter='2026-Q3')[:12])
    preview = workbook_io.parse_plan_workbook(make_plan(rows, headers))
    assert preview['valid'] and not preview['errors']
    assert preview['rows'][0]['month'] == '2026-09'
    assert preview['rows'][1]['voyage_id'] == 101
    assert all(row['week'] is None for row in preview['rows'])
    if len(headers) == 12:
        assert preview['rows'][2]['quarter'] == '2026-Q3'


@pytest.mark.parametrize('headers', [LEGACY_PLAN_COLUMNS_8, LEGACY_PLAN_COLUMNS_12])
def test_legacy_headers_cannot_hide_unlabelled_week_values(headers):
    preview = workbook_io.parse_plan_workbook(make_plan([month_row(period_type='week', month=None, week='2026-W38')], headers))
    assert not preview['valid'] and not preview['rows']
    assert preview['errors'][0]['row'] == 2
    assert 'ngoài các cột' in preview['errors'][0]['message']


@pytest.mark.parametrize('headers', [workbook_io.PLAN_COLUMNS, workbook_io.PLAN_KEYS])
def test_weekly_import_normalizes_iso_week_and_allows_future_and_year_crossing_targets(headers):
    weeks = [' 2026-W38 ', '2020-W53', '2099-W01']
    rows = [month_row(period_type='week', month=None, week=week) for week in weeks]
    preview = workbook_io.parse_plan_workbook(make_plan(rows, headers))
    assert preview['valid'] and not preview['errors']
    assert [row['week'] for row in preview['rows']] == ['2026-W38', '2020-W53', '2099-W01']
    assert all(row['period_type'] == 'week' and row['month'] is None for row in preview['rows'])


@pytest.mark.parametrize('changes', [
    {'week': None}, {'week': ''}, {'week': '2025-W53'}, {'week': '2026-W00'},
    {'week': '2026-W54'}, {'week': '2026-w38'}, {'week': '2026-W8'},
    {'week': '1999-W52'}, {'week': '2100-W01'}, {'week': 202638},
    {'week': '2026-W38', 'month': '2026-09'}, {'week': '2026-W38', 'start_date': '2026-09-14'},
])
def test_weekly_import_reports_invalid_or_conflicting_period_fields(changes):
    row = month_row(**{'period_type': 'week', 'month': None, **changes})
    preview = workbook_io.parse_plan_workbook(make_plan([row]))
    assert not preview['valid'] and not preview['rows']
    assert preview['errors'][0]['row'] == 2
    assert 'Kỳ kế hoạch' in preview['errors'][0]['message']


def test_weekly_import_detects_duplicate_after_normalizing_week_text():
    row = {'period_type': 'week', 'month': None, 'week': '2026-W38'}
    preview = workbook_io.parse_plan_workbook(make_plan([month_row(**row), month_row(**{**row, 'week': ' 2026-W38 '})]))
    assert not preview['valid'] and len(preview['rows']) == 1
    assert preview['errors'][0]['row'] == 3 and 'Trùng' in preview['errors'][0]['message']
    distinct = workbook_io.parse_plan_workbook(make_plan([
        month_row(**row), month_row(**row, terminal='ben_thuy'), month_row(**row, metric='teu'),
    ]))
    assert distinct['valid'] and len(distinct['rows']) == 3


def test_duplicate_plan_key_is_reported_without_silently_summing_or_replacing():
    preview = workbook_io.parse_plan_workbook(make_plan([month_row(amount=10), month_row(amount=20)]))
    assert not preview["valid"]
    assert len(preview["rows"]) == 1 and preview["rows"][0]["amount"] == "10"
    assert preview["errors"][0]["row"] == 3
    assert "Trùng" in preview["errors"][0]["message"]
    distinct = workbook_io.parse_plan_workbook(make_plan([month_row(), month_row(terminal="ben_thuy"), month_row(metric="teu")]))
    assert distinct["valid"] and len(distinct["rows"]) == 3


def test_file_size_and_invalid_zip_rejected_before_xml_loader(monkeypatch):
    monkeypatch.setattr(workbook_io, "load_workbook", lambda *args, **kwargs: pytest.fail("Rejected upload reached XML loader"))
    with pytest.raises(ValueError, match="2 MB"):
        workbook_io.parse_plan_workbook(b"0" * (workbook_io.MAX_UPLOAD_BYTES + 1))
    with pytest.raises(ValueError, match="xlsx"):
        workbook_io.parse_plan_workbook(b"not a workbook")


@pytest.mark.parametrize("case", ["inflated_size", "many_members", "macro", "external_link"])
def test_zip_limits_and_active_content_rejected_before_xml_loader(monkeypatch, case):
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        if case == "inflated_size":
            archive.writestr("xl/large.xml", b"0" * (20 * 1024 * 1024 + 1))
        elif case == "many_members":
            for index in range(1001):
                archive.writestr(f"small-{index}.xml", b"x")
        elif case == "macro":
            archive.writestr("xl/vbaProject.bin", b"not executed")
        else:
            archive.writestr("xl/externalLinks/externalLink1.xml", b"not followed")
    monkeypatch.setattr(workbook_io, "load_workbook", lambda *args, **kwargs: pytest.fail("Rejected ZIP reached XML loader"))
    with pytest.raises(ValueError):
        workbook_io.parse_plan_workbook(buffer.getvalue())


@pytest.mark.parametrize('headers,last_column', [(LEGACY_PLAN_COLUMNS_12, 'L'), (workbook_io.PLAN_COLUMNS, 'M')])
def test_row_count_and_forged_dimensions_are_bounded(headers, last_column):
    with pytest.raises(ValueError, match="500"):
        workbook_io.parse_plan_workbook(make_plan([month_row()] * 501))
    original = make_plan([month_row()[:len(headers)]], headers)
    expected = f'A1:{last_column}2'.encode()
    with zipfile.ZipFile(BytesIO(original)) as archive:
        assert expected in archive.read('xl/worksheets/sheet1.xml')
    forged = replace_zip_member(original, "xl/worksheets/sheet1.xml", lambda xml: xml.replace(expected, f'A1:{last_column}1048576'.encode()))
    with pytest.raises(ValueError, match="500"):
        workbook_io.parse_plan_workbook(forged)


def operation(index):
    return {"terminal_name": "Cảng mẫu", "id": f"{index:08d}", "operation_code": f"{index:010d}",
            "operation_date": "2026-09-01", "shift_code": "01", "vessel_name": "Tàu mẫu",
            "voyage_code": "0000000101", "cargo_name": "Hàng mẫu", "customer_name": "0012",
            "direction": "UNLOADING", "quantity": Decimal("1.25"), "quantity_unit_name": "M3",
            "weight": Decimal("2.125"), "weight_unit_name": "M3", "tonnage": None, "teu": 0}


def report(rows):
    return {"meta": {"report_id": "report-example", "filters": {"start_date": "2026-09-01", "end_date": "2026-09-09", "terminal": "all"},
                     "selection": {"customer_id": "000012", "voyage_id": "000101"}, "operations_filter": "all"},
            "overview": {"total_tonnage": None, "total_teu": 0, "record_count": len(rows), "tonnage_status": "unavailable", "teu_status": "complete"}}


def test_export_includes_full_row_count_and_preserves_types_native_units_and_nulls(tmp_path):
    rows = [operation(index) for index in range(1, 64)]
    rows[1]["weight"] = None
    rows[2]["weight"] = 0
    rows[3]["quantity"] = float("nan")
    target = tmp_path / "report-all-rows.xlsx"
    target.write_bytes(workbook_io.report_workbook(report(rows), rows))
    book = load_workbook(target, data_only=False)
    sheet = book["Tác nghiệp"]
    assert sheet.max_row == 64  # More than the dialog's 25-row page.
    assert sheet["B64"].value == "00000063"
    for coordinate, expected in {"B2": "00000001", "C2": "0000000001", "G2": "0000000101", "I2": "0012"}.items():
        assert sheet[coordinate].value == expected and sheet[coordinate].data_type == "s"
    assert isinstance(sheet["D2"].value, datetime) and sheet["D2"].value.date() == date(2026, 9, 1)
    assert sheet["K2"].value == 1.25 and sheet["K2"].data_type == "n"
    assert sheet["M2"].value == 2.125 and sheet["M2"].data_type == "n"
    assert sheet["N2"].value == "M3" and sheet["O2"].value is None
    assert sheet["M3"].value is None and sheet["M4"].value == 0
    assert sheet["K5"].value is None and sheet["P2"].value == 0
    summary = {row[0].value: row[1] for row in book["Tổng hợp"].iter_rows(min_row=2)}
    assert summary["Số dòng trong tệp"].value == 63
    assert summary["Số dòng nguồn"].value == 63
    assert summary["Sản lượng thông qua"].value is None
    assert summary["Container tác nghiệp"].value == 0
    assert summary["ID khách hàng"].value == "000012" and summary["ID khách hàng"].data_type == "s"
    book.close()


@pytest.mark.parametrize("text", ["=1+1", "+SUM(A1:A2)", "-1+2", "@SUM(A1:A2)", '=HYPERLINK("https://example.invalid","test")'])
def test_export_formula_like_names_and_identifiers_remain_literal_text(text):
    rows = [{**operation(1), "cargo_name": text, "customer_name": text, "operation_code": text}]
    content = workbook_io.report_workbook(report(rows), rows, title=text)
    book = load_workbook(BytesIO(content), data_only=False)
    for cell in (book["Tổng hợp"]["B2"], book["Tác nghiệp"]["C2"], book["Tác nghiệp"]["H2"], book["Tác nghiệp"]["I2"]):
        assert cell.value == text and cell.data_type == "s"
    book.close()
    with zipfile.ZipFile(BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                assert b"<f>" not in archive.read(name)


def test_shift_export_contains_all_rows_without_inventing_missing_metrics():
    rows = [operation(1)]
    shifts = [{"date": "2026-09-01", "shift_code": "001", "terminal_name": "Cảng mẫu", "tonnage": None, "teu": 0, "record_count": 1}]
    book = load_workbook(BytesIO(workbook_io.report_workbook(report(rows), rows, shifts=shifts)))
    sheet = book["Theo ca"]
    assert sheet.max_row == 2
    assert sheet["B2"].value == "001" and sheet["B2"].data_type == "s"
    assert sheet["D2"].value is None and sheet["E2"].value == 0 and sheet["F2"].value == 1
    book.close()


def test_numeric_metadata_identifiers_are_exported_as_text():
    snapshot = report([])
    snapshot["meta"]["selection"] = {"customer_id": 123, "voyage_id": 101}
    book = load_workbook(BytesIO(workbook_io.report_workbook(snapshot, [])))
    summary = {row[0].value: row[1] for row in book["Tổng hợp"].iter_rows(min_row=2)}
    assert summary["ID khách hàng"].value == "123" and summary["ID khách hàng"].data_type == "s"
    assert summary["ID chuyến"].value == "101" and summary["ID chuyến"].data_type == "s"
    book.close()


def test_export_preserves_scope_rule_and_first_berth_evidence_as_literal_values():
    from datetime import timezone
    rows = [{**operation(1), 'production_scope': 'vietsun', 'berth_assignment_status': 'assigned',
             'initial_berth_id': 13, 'initial_berth_code': '=Cau5',
             'initial_berth_at': datetime(2026, 8, 31, 9, 15, tzinfo=timezone.utc)}]
    snapshot = report(rows)
    snapshot['meta']['filters']['production_scope'] = 'vietsun'
    snapshot['meta']['berth_rule_version'] = 'initial-berth-v1'
    book = load_workbook(BytesIO(workbook_io.report_workbook(snapshot, rows)), data_only=False)
    summary = {row[0].value: row[1].value for row in book['Tổng hợp'].iter_rows(min_row=2)}
    assert summary['Phạm vi sản lượng'] == 'Cầu 5'
    assert summary['Mã phạm vi sản lượng'] == 'vietsun'
    assert summary['Phiên bản quy tắc cầu cập đầu tiên'] == 'initial-berth-v1'
    sheet = book['Tác nghiệp']
    assert sheet['Q2'].value == 'Cầu 5' and sheet['R2'].value == 'Đã xác định'
    assert sheet['S2'].value == '13' and sheet['S2'].data_type == 's'
    assert sheet['T2'].value == '=Cau5' and sheet['T2'].data_type == 's'
    assert sheet['U2'].value == '2026-08-31T09:15:00+00:00'
    book.close()


@pytest.mark.parametrize('status', ['missing', 'ambiguous'])
def test_unclassified_export_keeps_unknown_berth_empty(status):
    rows = [{**operation(1), 'production_scope': 'unclassified', 'berth_assignment_status': status,
             'initial_berth_id': None, 'initial_berth_code': None, 'initial_berth_at': None}]
    snapshot = report(rows)
    snapshot['meta']['filters']['production_scope'] = 'unclassified'
    snapshot['meta']['berth_rule_version'] = 'initial-berth-v1'
    book = load_workbook(BytesIO(workbook_io.report_workbook(snapshot, rows)))
    sheet = book['Tác nghiệp']
    assert sheet['Q2'].value == 'Chưa xác định cầu'
    assert sheet['R2'].value == {'missing': 'Thiếu dữ liệu cầu', 'ambiguous': 'Không xác định duy nhất'}[status]
    assert all(sheet.cell(2, col).value is None for col in (19, 20, 21))
    book.close()


def test_legacy_export_is_labelled_without_reclassifying_or_mutating_saved_report():
    from copy import deepcopy
    rows = [operation(1)]
    snapshot = report(rows)
    unchanged = deepcopy(snapshot)
    book = load_workbook(BytesIO(workbook_io.report_workbook(snapshot, rows)))
    summary = {row[0].value: row[1].value for row in book['Tổng hợp'].iter_rows(min_row=2)}
    assert summary['Phạm vi sản lượng'].startswith('Phạm vi cũ')
    assert summary['Mã phạm vi sản lượng'] == 'legacy'
    assert summary['Phiên bản quy tắc cầu cập đầu tiên'] == 'Chưa được lưu trong bản dữ liệu này'
    assert book['Tác nghiệp']['R2'].value == 'Chưa lưu bằng chứng phân loại'
    assert snapshot == unchanged
    book.close()


def test_weekly_progress_export_preserves_full_target_week_and_partial_actual():
    snapshot = report([])
    snapshot['meta']['filters'].update(start_date='2026-09-14', end_date='2026-09-18')
    snapshot['throughput_progress'] = {'items': [{
        'period_type': 'week', 'period_key': '2026-W38', 'start_date': '2026-09-14', 'end_date': '2026-09-20',
        'target': 1000, 'actual': 600, 'completion_percent': 60, 'provisional_completion_percent': None,
        'achieved': False, 'provisional': False, 'status': 'ready', 'target_source': 'company',
        'plans': [{'id': 17, 'version': 2, 'reference': 'KH-TUAN-38'}],
    }]}
    book = load_workbook(BytesIO(workbook_io.report_workbook(snapshot, [])))
    sheet = book['Mục tiêu thông qua']
    assert [sheet.cell(2, col).value for col in range(1, 8)] == [
        'Tuần', '2026-W38', '2026-09-14', '2026-09-20', 1000, 600, 60,
    ]
    assert sheet['L2'].value == '17 / v2' and sheet['M2'].value == 'KH-TUAN-38'
    book.close()


@pytest.mark.parametrize('complete', [True, False])
def test_export_preserves_captured_milestone_actuals_and_comparison_without_recalculation(complete):
    from copy import deepcopy
    snapshot = report([])
    snapshot['meta']['previous_period'] = {
        'mode': 'previous_year', 'label': 'Cùng ngày/tháng năm trước',
        'start_date': '2025-09-01', 'end_date': '2025-09-09',
    }
    snapshot['throughput_progress'] = {'items': [{
        'period_type': 'month', 'period_key': '2026-09', 'start_date': '2026-09-01', 'end_date': '2026-09-30',
        'target': 2000, 'actual': 1000, 'completion_percent': 50, 'provisional_completion_percent': None,
        'achieved': False, 'provisional': False, 'status': 'ready', 'target_source': 'company',
        'plans': [{'id': 17, 'version': 2, 'reference': 'KH-09'}],
        'pace': {'status': 'ready' if complete else 'unknown', 'milestone_date': '2026-09-05',
                 'target': 500, 'actual': 200 if complete else None,
                 'difference': -300 if complete else None, 'completion_percent': 40 if complete else None,
                 'reason': None if complete else 'Thiếu dữ liệu đến ngày mốc.'},
    }]}
    original = deepcopy(snapshot)
    book = load_workbook(BytesIO(workbook_io.report_workbook(snapshot, [])))
    summary = {row[0].value: row[1].value for row in book['Tổng hợp'].iter_rows(min_row=2)}
    assert summary['Cơ sở so sánh'] == 'Cùng ngày/tháng năm trước'
    assert summary['Kỳ so sánh từ ngày'] == '2025-09-01'
    assert summary['Kỳ so sánh đến ngày'] == '2025-09-09'
    assert book['Mục tiêu thông qua']['F2'].value == 1000
    pace = book['Tiến độ theo mốc']
    assert [pace.cell(2, column).value for column in range(1, 5)] == [
        '2026-09', '2026-09-01', '2026-09-05', 500,
    ]
    assert [pace.cell(2, column).value for column in range(5, 8)] == ([200, -300, 40] if complete else [None] * 3)
    assert pace['I2'].value == '17 / v2'
    if not complete:
        assert pace['H2'].value == 'Thiếu dữ liệu đến ngày mốc.'
    assert snapshot == original
    book.close()
